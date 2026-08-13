#!/usr/bin/env python3
"""Set the Status column in the send workbook, so seed batches are one command.

The flow only picks up rows whose Status is Ready. Holding everyone else back
is what turns the full list into a small test send, and doing it here rather
than by hand in Excel avoids a mistyped find-and-replace reaching 200 people.

Examples:
    # show where things stand
    python3 email/tools/set-status.py --show

    # seed batch: only these rows go out, everyone else is held
    python3 email/tools/set-status.py --only SHWC-0015,SHWC-0016,SHWC-0017

    # release everyone who is on Hold, leaving Sent and Unsubscribed alone
    python3 email/tools/set-status.py --release
"""

import argparse
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required:  pip3 install openpyxl")

DEFAULT_WB = Path("/Users/brianakumah/Documents/development/SHWC/SEND - 10 Year Anniversary.xlsx")
SHEET = "Contacts"

# Never touched by --only or --release. Once someone is Sent they must not be
# mailed twice, and an unsubscribe has to stick.
PROTECTED = {"Sent", "Unsubscribed", "Bounced"}


def columns(ws):
    header = [str(c.value).strip() if c.value else "" for c in ws[1]]
    need = ("ContactID", "FirstName", "LastName", "EmailAddress", "Status")
    missing = [n for n in need if n not in header]
    if missing:
        sys.exit(f"Workbook is missing column(s) {missing}")
    return {name: header.index(name) + 1 for name in header if name}


def show(ws, col):
    counts = Counter()
    for row in ws.iter_rows(min_row=2):
        if not row[col["ContactID"] - 1].value:
            continue
        counts[str(row[col["Status"] - 1].value or "").strip() or "(blank)"] += 1
    total = sum(counts.values())
    print(f"{total} rows")
    for status, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {status:14} {n:>4}")
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WB)
    parser.add_argument("--only", help="comma separated ContactIDs to leave Ready; all others go to Hold")
    parser.add_argument("--release", action="store_true", help="turn every Hold back into Ready")
    parser.add_argument("--show", action="store_true", help="print status counts and exit")
    args = parser.parse_args()

    if not args.workbook.exists():
        sys.exit(f"Workbook not found: {args.workbook}\nRun build-list.py first.")
    if not (args.only or args.release or args.show):
        parser.error("pick one of --only, --release, or --show")

    wb = load_workbook(args.workbook)
    if SHEET not in wb.sheetnames:
        sys.exit(f"No '{SHEET}' sheet in {args.workbook}")
    ws = wb[SHEET]
    col = columns(ws)

    if args.show:
        show(ws, col)
        print(f"\ntables present: {list(ws.tables)}")
        return

    wanted = set()
    if args.only:
        wanted = {p.strip().upper() for p in args.only.split(",") if p.strip()}

    changed, kept, protected, picked = 0, 0, 0, []
    for row in ws.iter_rows(min_row=2):
        cid_cell = row[col["ContactID"] - 1]
        if not cid_cell.value:
            continue
        cid = str(cid_cell.value).strip().upper()
        status_cell = row[col["Status"] - 1]
        status = str(status_cell.value or "").strip()

        if status in PROTECTED:
            protected += 1
            continue

        if args.release:
            if status == "Hold":
                status_cell.value = "Ready"
                changed += 1
            else:
                kept += 1
            continue

        target = "Ready" if cid in wanted else "Hold"
        if cid in wanted:
            picked.append((cid,
                           str(row[col["FirstName"] - 1].value or ""),
                           str(row[col["LastName"] - 1].value or ""),
                           str(row[col["EmailAddress"] - 1].value or "")))
        if status != target:
            status_cell.value = target
            changed += 1
        else:
            kept += 1

    missing = wanted - {p[0] for p in picked}
    if missing:
        sys.exit(f"These ContactIDs are not in the workbook: {sorted(missing)}\n"
                 "Nothing was saved.")

    wb.save(args.workbook)

    if picked:
        print("Ready to send:")
        for cid, first, last, email in picked:
            print(f"  {cid}  {first} {last} <{email}>")
        print()
    print(f"changed {changed}, unchanged {kept}, protected {protected}")
    print()
    wb2 = load_workbook(args.workbook)
    show(wb2[SHEET], col)
    print(f"\ntables present: {list(wb2[SHEET].tables)}")


if __name__ == "__main__":
    main()
