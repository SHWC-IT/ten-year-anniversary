#!/usr/bin/env python3
"""Render the anniversary invitation with merge tokens filled in.

This mirrors exactly what the Power Automate flow does to the template: a
plain string replace of each {{Token}}. Running it proves the substitution
works, and produces a file you can open in a browser to check the design,
before a single real email is sent.

Examples:
    # sample name, writes email/anniversary-invite.preview.html
    python3 email/tools/render-preview.py

    # a specific person, pulled from the generated send workbook
    python3 email/tools/render-preview.py --contact-id SHWC-0105 --out /tmp/one.html

    # just audit the tokens, write nothing
    python3 email/tools/render-preview.py --check

Careful with --contact-id: the default output file is committed to a public
repo and published with the site, so never render a real member's details into
it. Pass --out to somewhere outside the repo when spot-checking real rows.
"""

import argparse
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
EMAIL_DIR = HERE.parent
TEMPLATE = EMAIL_DIR / "anniversary-invite.html"
PREVIEW = EMAIL_DIR / "anniversary-invite.preview.html"
WORKBOOK = Path("/Users/brianakumah/Documents/development/SHWC/SEND - 10 Year Anniversary.xlsx")

# The tokens the flow knows how to replace. Anything else in the template is a
# mistake, so we fail loudly rather than mailing a literal {{Token}}.
KNOWN_TOKENS = {"FirstName", "EmailAddress"}

TOKEN_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")


def tokens_in(text):
    return sorted(set(TOKEN_RE.findall(text)))


def lookup_contact(contact_id):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is required to read the workbook:  pip3 install openpyxl")
    if not WORKBOOK.exists():
        sys.exit(f"Workbook not found, run build-list.py first: {WORKBOOK}")
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Contacts"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(c) if c is not None else "" for c in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    for row in rows[1:]:
        if not row or not row[idx["ContactID"]]:
            continue
        if str(row[idx["ContactID"]]).strip().lower() == contact_id.strip().lower():
            return (str(row[idx["FirstName"]] or ""),
                    str(row[idx["EmailAddress"]] or ""))
    sys.exit(f"No contact with ContactID {contact_id}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    # Placeholder by default: this file ships in a public repo.
    parser.add_argument("--name", default="Friend", help="value for {{FirstName}}")
    parser.add_argument("--email", default="member@example.com",
                        help="value for {{EmailAddress}}")
    parser.add_argument("--contact-id", help="pull name and email from the send workbook")
    parser.add_argument("--template", type=Path, default=TEMPLATE)
    parser.add_argument("--out", type=Path, default=PREVIEW)
    parser.add_argument("--check", action="store_true",
                        help="audit tokens only, write no file")
    args = parser.parse_args()

    if not args.template.exists():
        sys.exit(f"Template not found: {args.template}")
    html = args.template.read_text(encoding="utf-8")

    found = tokens_in(html)
    unknown = [t for t in found if t not in KNOWN_TOKENS]
    print(f"template        {args.template.name}")
    print(f"tokens found    {found or 'none'}")
    if unknown:
        sys.exit(f"Unknown token(s) {unknown}. The flow only replaces {sorted(KNOWN_TOKENS)}, "
                 "so these would be mailed out literally.")
    missing = sorted(KNOWN_TOKENS - set(found))
    if missing:
        print(f"note            declared but unused: {missing}")

    if args.check:
        print("check           OK")
        return

    if args.contact_id:
        name, email = lookup_contact(args.contact_id)
        print(f"contact         {args.contact_id}  {name}  {email}")
    else:
        name, email = args.name, args.email

    rendered = html.replace("{{FirstName}}", name).replace("{{EmailAddress}}", email)

    leftover = tokens_in(rendered)
    if leftover:
        sys.exit(f"Substitution failed, tokens remain: {leftover}")

    args.out.write_text(rendered, encoding="utf-8")
    print(f"rendered as     {name} <{email}>")
    print(f"output          {args.out}")
    print("tokens remaining none")


if __name__ == "__main__":
    main()
